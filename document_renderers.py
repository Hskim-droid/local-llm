"""Render and reopen the three user-selectable office output formats.

The renderer contract is intentionally small: a title/source line, one table
of normalized records, and evidence references. Each renderer writes a new
artifact and has a matching reopen check; no renderer writes back to the source
UI or sends the artifact.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping, Sequence


FORMATS = ("docx", "xlsx", "pptx")


class RendererError(RuntimeError):
    """A requested artifact cannot be rendered or reopened safely."""


def _validate_inputs(
    output_path: Path,
    output_format: str,
    columns: Sequence[str],
    records: Sequence[Mapping[str, Any]],
) -> tuple[str, ...]:
    if output_format not in FORMATS:
        raise ValueError(f"unsupported output format: {output_format}")
    expected_suffix = f".{output_format}"
    if output_path.suffix.lower() != expected_suffix:
        raise ValueError(f"output path must use {expected_suffix}")
    normalized_columns = tuple(columns)
    if not normalized_columns or len(set(normalized_columns)) != len(normalized_columns):
        raise ValueError("columns must contain unique names")
    for index, record in enumerate(records):
        if any(column not in record for column in normalized_columns):
            raise ValueError(f"record {index + 1} is missing a document column")
    return normalized_columns


def _cell_values(columns: Sequence[str], records: Sequence[Mapping[str, Any]]) -> list[list[str]]:
    return [[str(record[column]) for column in columns] for record in records]


def _render_docx(
    output_path: Path,
    source_system: str,
    captured_at: str,
    columns: tuple[str, ...],
    records: Sequence[Mapping[str, Any]],
) -> None:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - setup failure path
        raise RendererError("python-docx is required for DOCX output") from exc
    document = Document()
    document.add_heading("QMS Daily Draft", level=1)
    document.add_paragraph(f"Source: {source_system} | Captured: {captured_at}")
    document.add_paragraph("Draft only. No ERP/QMS write or external send was performed.")
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, columns):
        cell.text = heading
    for values in _cell_values(columns, records):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    document.add_heading("Evidence", level=2)
    for record in records:
        document.add_paragraph(f"{record['record_id']}: {record['source_ref']}")
    document.save(output_path)


def _render_xlsx(
    output_path: Path,
    source_system: str,
    captured_at: str,
    columns: tuple[str, ...],
    records: Sequence[Mapping[str, Any]],
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - setup failure path
        raise RendererError("openpyxl is required for XLSX output") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    sheet.cell(1, 1, "QMS Daily Draft").font = Font(bold=True, size=14)
    sheet.cell(2, 1, f"Source: {source_system} | Captured: {captured_at}")
    for column_index, heading in enumerate(columns, 1):
        cell = sheet.cell(4, column_index, heading)
        cell.data_type = "s"
        cell.font = Font(bold=True)
    for row_index, values in enumerate(_cell_values(columns, records), 5):
        for column_index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column_index, value)
            cell.data_type = "s"
    for column_index, heading in enumerate(columns, 1):
        longest = max([len(str(heading))] + [len(str(record[heading])) for record in records])
        sheet.column_dimensions[get_column_letter(column_index)].width = min(40, max(12, longest + 2))
    evidence = workbook.create_sheet("Evidence")
    evidence.append(("record_id", "source_ref", "evidence_selector"))
    for record in records:
        evidence.append((record.get("record_id", ""), record.get("source_ref", ""), record.get("evidence_selector", "")))
    for row in evidence.iter_rows():
        for cell in row:
            cell.data_type = "s"
    workbook.save(output_path)


def _render_pptx(
    output_path: Path,
    source_system: str,
    captured_at: str,
    columns: tuple[str, ...],
    records: Sequence[Mapping[str, Any]],
) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover - setup failure path
        raise RendererError("python-pptx is required for PPTX output") from exc
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "QMS Daily Draft"
    subtitle = title_slide.placeholders[1]
    subtitle.text = f"Source: {source_system} | Captured: {captured_at}"
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = table_slide.shapes.add_textbox(Inches(0.35), Inches(0.2), Inches(9.3), Inches(0.45))
    title_box.text_frame.text = "Records"
    title_box.text_frame.paragraphs[0].font.size = Pt(22)
    table = table_slide.shapes.add_table(len(records) + 1, len(columns), Inches(0.35), Inches(0.85), Inches(9.3), Inches(4.8)).table
    for column_index, heading in enumerate(columns):
        table.cell(0, column_index).text = heading
    for row_index, values in enumerate(_cell_values(columns, records), 1):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    evidence_box = table_slide.shapes.add_textbox(Inches(0.35), Inches(5.85), Inches(9.3), Inches(1.0))
    evidence_box.text_frame.text = "Evidence: " + "; ".join(
        f"{record.get('record_id', '')} → {record.get('source_ref', '')}" for record in records
    )
    presentation.save(output_path)


def render_artifact(
    output_path: Path,
    output_format: str,
    source_system: str,
    captured_at: str,
    columns: Sequence[str],
    records: Sequence[Mapping[str, Any]],
) -> None:
    normalized_columns = _validate_inputs(output_path, output_format, columns, records)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_format == "docx":
        _render_docx(output_path, source_system, captured_at, normalized_columns, records)
    elif output_format == "xlsx":
        _render_xlsx(output_path, source_system, captured_at, normalized_columns, records)
    else:
        _render_pptx(output_path, source_system, captured_at, normalized_columns, records)


def _reopen_docx(output_path: Path, columns: tuple[str, ...], records: Sequence[Mapping[str, Any]]) -> bool:
    try:
        from docx import Document
        document = Document(output_path)
        headers = [cell.text.strip() for cell in document.tables[0].rows[0].cells]
        rows = [[cell.text.strip() for cell in row.cells] for row in document.tables[0].rows[1:]]
    except (ImportError, IndexError, OSError) as exc:
        raise RendererError(f"cannot reopen DOCX artifact: {exc}") from exc
    return bool(document.paragraphs) and tuple(headers) == columns and rows == _cell_values(columns, records)


def _reopen_xlsx(output_path: Path, columns: tuple[str, ...], records: Sequence[Mapping[str, Any]]) -> bool:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(output_path, read_only=True, data_only=True)
        sheet = workbook["Report"]
        rows_with_headers = [list(row) for row in sheet.iter_rows(min_row=4, values_only=True)]
        headers = rows_with_headers[0]
        rows = rows_with_headers[1:]
        max_column = sheet.max_column
        max_row = sheet.max_row
        workbook.close()
    except (ImportError, KeyError, OSError, StopIteration) as exc:
        raise RendererError(f"cannot reopen XLSX artifact: {exc}") from exc
    return (
        max_column == len(columns)
        and max_row == 4 + len(records)
        and tuple(headers) == columns
        and rows == _cell_values(columns, records)
    )


def _reopen_pptx(output_path: Path, columns: tuple[str, ...], records: Sequence[Mapping[str, Any]]) -> bool:
    try:
        from pptx import Presentation
        presentation = Presentation(output_path)
        tables = [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]
        if len(tables) != 1:
            return False
        table = tables[0]
        if len(table.rows) != len(records) + 1 or len(table.columns) != len(columns):
            return False
        headers = [table.cell(0, column_index).text.strip() for column_index in range(len(columns))]
        rows = [
            [table.cell(row_index, column_index).text.strip() for column_index in range(len(columns))]
            for row_index in range(1, len(records) + 1)
        ]
    except (ImportError, IndexError, OSError) as exc:
        raise RendererError(f"cannot reopen PPTX artifact: {exc}") from exc
    return tuple(headers) == columns and rows == _cell_values(columns, records)


def reopen_artifact(
    output_path: Path,
    output_format: str,
    columns: Sequence[str],
    records: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    normalized_columns = _validate_inputs(output_path, output_format, columns, records)
    if output_format == "docx":
        passed = _reopen_docx(output_path, normalized_columns, records)
    elif output_format == "xlsx":
        passed = _reopen_xlsx(output_path, normalized_columns, records)
    else:
        passed = _reopen_pptx(output_path, normalized_columns, records)
    return {"name": f"{output_format}_reopen_and_cell_values", "passed": passed}
