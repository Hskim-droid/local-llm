import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from document_renderers import render_artifact, reopen_artifact  # noqa: E402


HAS_OFFICE_DEPS = all(importlib.util.find_spec(name) is not None for name in ("docx", "openpyxl", "pptx"))


@unittest.skipUnless(HAS_OFFICE_DEPS, "office renderer dependencies are not installed")
class DocumentRendererTests(unittest.TestCase):
    columns = ("record_id", "title", "status", "owner", "updated_at")
    records = [
        {
            "record_id": "QMS-1",
            "title": "=SUM(A1)",
            "status": "Open",
            "owner": "J. Kim",
            "updated_at": "2026-09-21T09:00:00+00:00",
            "source_ref": "qms://issue/QMS-1",
        }
    ]

    def test_all_formats_reopen_and_preserve_formula_like_source_as_text(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for output_format in ("docx", "xlsx", "pptx"):
                path = root / f"artifact.{output_format}"
                render_artifact(path, output_format, "QMS", "2026-09-21T09:00:00+00:00", self.columns, self.records)
                self.assertTrue(reopen_artifact(path, output_format, self.columns, self.records)["passed"])
            from openpyxl import load_workbook

            workbook = load_workbook(root / "artifact.xlsx", read_only=False, data_only=False)
            cell = workbook["Report"]["B5"]
            self.assertEqual(cell.value, "=SUM(A1)")
            self.assertEqual(cell.data_type, "s")
            workbook.close()

    def test_output_suffix_is_part_of_the_format_contract(self):
        with tempfile.TemporaryDirectory() as temp:
            with self.assertRaises(ValueError):
                render_artifact(Path(temp) / "artifact.docx", "xlsx", "QMS", "now", self.columns, self.records)

    def test_reopen_rejects_extra_rows_in_xlsx_and_pptx(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            xlsx_path = root / "extra.xlsx"
            render_artifact(xlsx_path, "xlsx", "QMS", "now", self.columns, self.records)
            from openpyxl import load_workbook

            workbook = load_workbook(xlsx_path)
            workbook["Report"].cell(6, 1, "EXTRA")
            workbook.save(xlsx_path)
            workbook.close()
            self.assertFalse(reopen_artifact(xlsx_path, "xlsx", self.columns, self.records)["passed"])

            pptx_path = root / "extra.pptx"
            render_artifact(pptx_path, "pptx", "QMS", "now", self.columns, self.records)
            from pptx import Presentation
            from pptx.util import Inches

            presentation = Presentation(pptx_path)
            presentation.slides[1].shapes.add_table(1, len(self.columns), Inches(0.35), Inches(6.9), Inches(9.3), Inches(0.4))
            presentation.save(pptx_path)
            self.assertFalse(reopen_artifact(pptx_path, "pptx", self.columns, self.records)["passed"])


if __name__ == "__main__":
    unittest.main()
