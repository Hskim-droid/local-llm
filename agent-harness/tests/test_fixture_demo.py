import json
import tempfile
import unittest
from unittest import mock
from pathlib import Path

import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import fixture_demo  # noqa: E402
import harness  # noqa: E402
from translation_pipeline import PassthroughTranslator  # noqa: E402


@unittest.skipUnless(fixture_demo.IMPORT_ERROR is None, "fixture dependencies are not installed")
class FixtureDemoTests(unittest.TestCase):
    def test_synthetic_qms_screen_becomes_reopenable_docx(self):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            config_path = temp_path / "config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "persona_id": "han-gyeol",
                        "persona_name": "한결",
                        "mode": "draft_only",
                        "state_dir": str(temp_path / "state"),
                        "artifact_dir": str(temp_path / "artifacts"),
                        "default_output_format": "docx",
                        "recipient_allowlist": [],
                        "executor": {"extract": "fixture-demo", "render": None, "send": None},
                    }
                ),
                encoding="utf-8",
            )
            self.assertEqual(
                harness.main(
                    [
                        "--config",
                        str(config_path),
                        "enqueue",
                        "--source-system",
                        "QMS",
                        "--request",
                        "synthetic daily report 2026-09-21",
                        "--output-format",
                        "docx",
                    ]
                ),
                0,
            )
            config = harness.load_config(config_path)
            with harness.connect(config) as connection:
                job_id = connection.execute("SELECT id FROM jobs").fetchone()[0]
            result = fixture_demo.run_demo(
                config_path,
                job_id,
                ROOT / "fixtures" / "qms_daily.html",
                claim=True,
                translation_target="en",
                translation_backend="passthrough",
            )
            self.assertEqual(result["records"], 3)
            self.assertTrue(result["translated"])
            self.assertTrue(Path(result["artifact"]).is_file())
            self.assertTrue(Path(result["manifest"]).is_file())
            manifest = json.loads(Path(result["manifest"]).read_text(encoding="utf-8"))
            self.assertEqual(manifest["translation"]["target_language"], "en")
            self.assertTrue(manifest["translation"]["local_transport_only"])
            self.assertEqual(manifest["translation"]["transport_scope"], "in-process")
            with harness.connect(config) as connection:
                row = connection.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()
            self.assertEqual(row["status"], "drafted")

    def test_translation_stage_keeps_source_and_puts_changed_text_in_docx(self):
        class PrefixTranslator(PassthroughTranslator):
            backend_name = "deterministic-fixture"
            model_name = "fixture-model"

            def translate_text(self, text, request):
                return f"[EN] {text}"

        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            config_path = temp_path / "config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "persona_id": "han-gyeol",
                        "persona_name": "한결",
                        "mode": "draft_only",
                        "state_dir": str(temp_path / "state"),
                        "artifact_dir": str(temp_path / "artifacts"),
                        "default_output_format": "docx",
                        "recipient_allowlist": [],
                        "executor": {"extract": "fixture-demo", "render": None, "send": None},
                    }
                ),
                encoding="utf-8",
            )
            self.assertEqual(
                harness.main(
                    [
                        "--config",
                        str(config_path),
                        "enqueue",
                        "--source-system",
                        "QMS",
                        "--request",
                        "synthetic translated report",
                        "--output-format",
                        "docx",
                    ]
                ),
                0,
            )
            config = harness.load_config(config_path)
            with harness.connect(config) as connection:
                job_id = connection.execute("SELECT id FROM jobs").fetchone()[0]
            with mock.patch("fixture_demo.build_translator", return_value=PrefixTranslator()):
                result = fixture_demo.run_demo(
                    config_path,
                    job_id,
                    ROOT / "fixtures" / "qms_daily.html",
                    claim=True,
                    translation_target="en",
                    translation_backend="passthrough",
                )
            manifest = json.loads(Path(result["manifest"]).read_text(encoding="utf-8"))
            item = manifest["translation"]["items"][0]
            self.assertNotEqual(item["fields"]["title"]["source"], item["fields"]["title"]["translated"])
            from docx import Document

            document = Document(result["artifact"])
            self.assertIn("[EN] Incoming inspection hold", document.tables[0].rows[1].cells[1].text)

    def test_duplicate_screen_rows_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            html = Path(temp) / "duplicate.html"
            fixture = (ROOT / "fixtures" / "qms_daily.html").read_text(encoding="utf-8")
            duplicate = fixture.replace(
                '<tr data-record-id="QMS-1003">\n            <td>QMS-1003</td>',
                '<tr data-record-id="QMS-1001">\n            <td>QMS-1001</td>',
            )
            html.write_text(duplicate, encoding="utf-8")
            with self.assertRaises(ValueError):
                fixture_demo.extract_qms_screen(html)

    def test_fixture_renders_requested_xlsx_and_pptx(self):
        class PrefixTranslator(PassthroughTranslator):
            backend_name = "deterministic-fixture"
            model_name = "fixture-model"

            def translate_text(self, text, request):
                return f"[EN] {text}"

        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            config_path = temp_path / "config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "persona_id": "han-gyeol",
                        "persona_name": "한결",
                        "mode": "draft_only",
                        "state_dir": str(temp_path / "state"),
                        "artifact_dir": str(temp_path / "artifacts"),
                        "default_output_format": "docx",
                        "recipient_allowlist": [],
                        "executor": {"extract": "fixture-demo", "render": None, "send": None},
                    }
                ),
                encoding="utf-8",
            )
            for output_format in ("xlsx", "pptx"):
                self.assertEqual(
                    harness.main(
                        [
                            "--config",
                            str(config_path),
                            "enqueue",
                            "--source-system",
                            "QMS",
                            "--request",
                            f"synthetic {output_format} report 2026-09-21",
                            "--output-format",
                            output_format,
                        ]
                    ),
                    0,
                )
                config = harness.load_config(config_path)
                with harness.connect(config) as connection:
                    job_id = connection.execute("SELECT id FROM jobs WHERE output_format=?", (output_format,)).fetchone()[0]
                with mock.patch("fixture_demo.build_translator", return_value=PrefixTranslator()):
                    result = fixture_demo.run_demo(
                        config_path,
                        job_id,
                        ROOT / "fixtures" / "qms_daily.html",
                        claim=True,
                        translation_target="en",
                        translation_backend="passthrough",
                    )
                self.assertEqual(result["format"], output_format)
                self.assertTrue(Path(result["artifact"]).is_file())
                self.assertEqual(Path(result["artifact"]).suffix, f".{output_format}")
                manifest = json.loads(Path(result["manifest"]).read_text(encoding="utf-8"))
                self.assertNotEqual(
                    manifest["translation"]["items"][0]["fields"]["title"]["source"],
                    manifest["translation"]["items"][0]["fields"]["title"]["translated"],
                )
                if output_format == "xlsx":
                    from openpyxl import load_workbook

                    workbook = load_workbook(result["artifact"], read_only=True)
                    self.assertIn("Evidence", workbook.sheetnames)
                    self.assertEqual(workbook["Report"]["B5"].value, "[EN] Incoming inspection hold")
                    workbook.close()
                else:
                    from pptx import Presentation

                    presentation = Presentation(result["artifact"])
                    self.assertGreaterEqual(len(presentation.slides), 2)
                    table = next(shape.table for shape in presentation.slides[1].shapes if shape.has_table)
                    self.assertEqual(table.cell(1, 1).text, "[EN] Incoming inspection hold")
                with harness.connect(config) as connection:
                    status = connection.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()[0]
                self.assertEqual(status, "drafted")


if __name__ == "__main__":
    unittest.main()
